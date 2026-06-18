"""Document parse service.

Responsible for the parse stage only: file → normalized chunks + image entries.
Does NOT perform image captioning, chunk building, or embedding.

The result is a ``ParseResult`` dataclass containing:
- ``chunks_with_metadata`` — list of chunk dicts with text, page metadata, and
  image/table references (no captions attached yet).
- ``images_data`` — raw image entries: ``{path, page_number, mime_type}``.
- ``parse_elapsed_s`` — wall-clock seconds for the parse operation.
- ``backend_used`` — e.g. ``"mineru/pipeline"``, ``"excel"``, ``"text"``.
"""

from __future__ import annotations

import json
import logging
import mimetypes
import os
import re
import subprocess
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from langchain_community.document_loaders import TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import Settings
from app.services.document_chunk_builder import DocumentChunkBuilder

logger = logging.getLogger(__name__)


@dataclass
class ParseResult:
    chunks_with_metadata: list[dict[str, Any]]
    images_data: list[dict[str, Any]]
    parse_elapsed_s: float
    backend_used: str


class DocumentParseService:
    """Parse-only service: file path + document metadata → ParseResult.

    Does not perform image captioning, chunk building for indexing, or embedding.
    Thread-safe when wired as a Factory (fresh instance per task).
    """

    EXCEL_EXTENSIONS: frozenset[str] = frozenset({".xlsx"})
    MINERU_EXTENSIONS: frozenset[str] = frozenset({".pdf", ".docx", ".pptx", ".html", ".md"})

    def __init__(self, settings: Settings, chunk_builder: DocumentChunkBuilder | None = None):
        self.settings = settings
        self.document_chunk_builder = chunk_builder or DocumentChunkBuilder(
            target_tokens=settings.rag_chunk_target_tokens,
            overlap_tokens=settings.rag_chunk_overlap_tokens,
            max_tokens=settings.rag_chunk_max_tokens,
        )
        self._mineru_output_path: str | None = None

    async def parse_document(
        self,
        file_path: str,
        filename: str,
        document_id: str,
    ) -> ParseResult:
        """Parse a document file and return normalized chunks + image entries.

        Does NOT perform image captioning or embedding.

        Dispatches by extension:
        - ``.txt`` → plain-text loader
        - ``.xlsx`` → openpyxl workbook parser
        - ``.pdf``, ``.docx``, ``.pptx``, ``.html``, ``.md`` → MinerU pipeline
        """
        start_time = time.perf_counter()
        ext = os.path.splitext(filename)[1].lower()

        if ext == ".txt":
            loader = TextLoader(file_path, encoding="utf-8")
            documents = loader.load()
            chunks = self._create_chunks(documents)
            chunks_with_metadata = [{"text": chunk} for chunk in chunks]
            images_data: list[dict[str, Any]] = []
            backend_used = "text"

        elif ext in self.EXCEL_EXTENSIONS:
            chunks_with_metadata = self._process_excel_workbook(file_path, filename)
            images_data = []
            backend_used = "excel"

        elif ext in self.MINERU_EXTENSIONS:
            chunks_with_metadata, images_data = await self._process_with_mineru(
                file_path, document_id, filename
            )
            backend_raw = str(getattr(self.settings, "mineru_backend", "pipeline") or "pipeline")
            backend_used = f"mineru/{backend_raw.strip().lower()}"

        else:
            raise ValueError(f"Unsupported file type: {filename}")

        parse_elapsed_s = time.perf_counter() - start_time
        return ParseResult(
            chunks_with_metadata=chunks_with_metadata,
            images_data=images_data,
            parse_elapsed_s=parse_elapsed_s,
            backend_used=backend_used,
        )

    async def _process_with_mineru(
        self,
        file_path: str,
        document_id: str,
        original_filename: str | None = None,
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Run MinerU on a document and return (chunks_with_metadata, images_data).

        The output directory is named ``mineru_output_{document_id}`` so that
        concurrent uploads never share temp paths.
        """
        try:
            temp_dir = Path(self.settings.temp_storage_path)
            output_dir = temp_dir / f"mineru_output_{document_id}"
            output_dir.mkdir(parents=True, exist_ok=True)
            self._mineru_output_path = str(output_dir)

            backend_raw = (
                str(getattr(self.settings, "mineru_backend", "pipeline") or "pipeline")
                .strip()
                .lower()
            )
            backend_aliases = {
                "vlm": "vlm-auto-engine",
                "hybrid": "hybrid-auto-engine",
            }
            backend = backend_aliases.get(backend_raw, backend_raw)
            valid_backends = {
                "pipeline",
                "hybrid-auto-engine",
                "hybrid-http-client",
                "vlm-auto-engine",
                "vlm-http-client",
            }
            if backend not in valid_backends:
                raise ValueError(
                    "Invalid MinerU backend "
                    f"'{backend_raw}'. Allowed values: {', '.join(sorted(valid_backends))}"
                )

            extra_args: list[str] = list(getattr(self.settings, "mineru_extra_args", []) or [])
            api_url = str(getattr(self.settings, "mineru_api_url", "") or "").strip()

            formula_flag = (
                "true" if getattr(self.settings, "extract_formulas_from_pdf", True) else "false"
            )
            table_flag = (
                "true" if getattr(self.settings, "extract_tables_from_pdf", True) else "false"
            )

            method = (
                str(getattr(self.settings, "mineru_method", "auto") or "auto").strip().lower()
            )
            valid_methods = {"auto", "txt", "ocr"}
            if method not in valid_methods:
                raise ValueError(
                    f"Invalid MinerU method '{method}'. Allowed values: {', '.join(sorted(valid_methods))}"
                )

            lang = str(getattr(self.settings, "mineru_lang", "") or "").strip()
            supports_method_and_lang = backend == "pipeline" or backend.startswith("hybrid-")

            cmd = [
                "mineru",
                "-p",
                file_path,
                "-o",
                str(output_dir),
                "--backend",
                backend,
                "-f",
                formula_flag,
                "-t",
                table_flag,
            ]

            if supports_method_and_lang:
                cmd.extend(["-m", method])
                if lang:
                    cmd.extend(["-l", lang])

            if api_url:
                cmd.extend(["--api-url", api_url])

            cmd.extend(extra_args)

            logger.info(
                "Running MinerU (backend=%s, method=%s, api_url=%s) for document %s: %s",
                backend,
                method if supports_method_and_lang else "n/a",
                "configured" if api_url else "auto-local",
                document_id,
                " ".join(cmd),
            )

            parse_started_at = time.perf_counter()
            result = subprocess.run(
                cmd,
                timeout=self.settings.mineru_timeout,
                check=True,
                capture_output=True,
                text=True,
            )
            parse_elapsed = time.perf_counter() - parse_started_at
            logger.info(
                "MinerU completed for %s in %.2fs (backend=%s, method=%s, api_url=%s)",
                document_id,
                parse_elapsed,
                backend,
                method if supports_method_and_lang else "n/a",
                "configured" if api_url else "auto-local",
            )

            if result.stdout:
                logger.debug("MinerU stdout for %s:\n%s", document_id, result.stdout)

            filename_without_ext = Path(file_path).stem
            filename_aliases = self._build_filename_aliases(filename_without_ext, original_filename)
            base_output_dir = self._resolve_mineru_output_dir(output_dir, filename_aliases)

            search_roots: list[Path] = []
            if base_output_dir is not None and base_output_dir.exists():
                search_roots.append(base_output_dir)
            else:
                logger.warning(
                    "MinerU output directory missing for %s. Searching entire output tree.",
                    filename_without_ext,
                )

            if output_dir not in search_roots:
                search_roots.append(output_dir)

            markdown_file = self._resolve_markdown_file(search_roots, filename_aliases)
            images_dir = markdown_file.parent / "images"

            content_list_path = (
                markdown_file.parent / f"{markdown_file.stem}_content_list.json"
            )
            content_blocks = None

            if content_list_path.exists():
                try:
                    content_blocks = self._parse_content_list_json(content_list_path)
                    logger.info(
                        "Loaded %d content blocks from content_list.json for %s",
                        len(content_blocks),
                        original_filename or filename_without_ext,
                    )
                except Exception as e:
                    logger.warning(
                        "Failed to parse content_list.json for %s: %s. Falling back to markdown.",
                        original_filename or filename_without_ext,
                        str(e),
                    )
                    content_blocks = None
            else:
                logger.info(
                    "content_list.json not found for %s, using markdown fallback",
                    original_filename or filename_without_ext,
                )

            images_by_path: dict[str, int] = {}

            if content_blocks:
                for block in content_blocks:
                    block_type = block.get("type")
                    if block_type not in ("image", "table"):
                        continue
                    img_path = block.get("img_path", "")
                    page_idx = block.get("page_idx")
                    if img_path and page_idx is not None:
                        images_by_path[img_path] = page_idx

            images_data: list[dict[str, Any]] = []
            page_to_images: dict[int, list[dict[str, Any]]] = {}

            if images_dir.exists():
                for img_file in images_dir.iterdir():
                    if img_file.is_file() and img_file.suffix.lower() in [
                        ".png",
                        ".jpg",
                        ".jpeg",
                    ]:
                        relative_path = f"images/{img_file.name}"
                        page_number = images_by_path.get(relative_path)

                        mime_type, _ = mimetypes.guess_type(str(img_file))
                        if not mime_type:
                            suffix = img_file.suffix.lower()
                            if suffix in {".jpg", ".jpeg"}:
                                mime_type = "image/jpeg"
                            elif suffix == ".png":
                                mime_type = "image/png"
                            elif suffix == ".gif":
                                mime_type = "image/gif"
                        mime_type = mime_type or f"image/{img_file.suffix.lstrip('.').lower()}"

                        image_entry: dict[str, Any] = {
                            "path": str(img_file),
                            "page_number": page_number,
                            "mime_type": mime_type,
                        }
                        images_data.append(image_entry)

                        if page_number is not None:
                            page_to_images.setdefault(page_number, []).append(image_entry)

            if content_blocks:
                chunks_with_metadata = self._create_chunks_with_page_metadata(
                    content_blocks,
                    page_to_images,
                    max_chunk_size=self._legacy_char_chunk_size(),
                )
            else:
                with open(markdown_file, encoding="utf-8") as f:
                    markdown_content = f.read()

                documents = [
                    type(
                        "Document",
                        (),
                        {"page_content": markdown_content, "metadata": {}},
                    )()
                ]
                chunks = self._create_chunks(documents)

                unpaged_images = [img for img in images_data if img["page_number"] is None]

                chunks_with_metadata = []
                for chunk in chunks:
                    chunks_with_metadata.append(
                        {
                            "text": chunk,
                            "page_start": None,
                            "page_end": None,
                            "has_images": bool(unpaged_images),
                            "image_count": len(unpaged_images),
                            "has_tables": False,
                            "table_count": 0,
                        }
                    )

            return chunks_with_metadata, images_data

        except subprocess.TimeoutExpired as exc:
            logger.error(
                "MinerU timed out after %ss while processing %s",
                self.settings.mineru_timeout,
                file_path,
            )
            raise RuntimeError(
                f"MinerU timed out after {self.settings.mineru_timeout}s"
            ) from exc
        except subprocess.CalledProcessError as exc:
            combined = "\n".join(filter(None, [exc.stdout, exc.stderr]))
            backend_for_log = locals().get(
                "backend", getattr(self.settings, "mineru_backend", "pipeline")
            )
            logger.error(
                "MinerU (backend=%s) failed (exit %s) while processing %s:\n%s",
                backend_for_log,
                exc.returncode,
                file_path,
                combined or "(no output captured)",
            )
            raise RuntimeError(f"MinerU failed with error: {combined}") from exc
        except Exception as exc:
            logger.error(
                "Unexpected MinerU error while processing %s: %s", file_path, exc
            )
            raise RuntimeError(
                f"Unexpected error in MinerU processing: {str(exc)}"
            ) from exc

    def _legacy_char_chunk_size(self) -> int:
        """Approximate char count for a target-token chunk (~4 chars/token heuristic)."""
        return max(200, int(self.settings.rag_chunk_target_tokens * 4))

    def _legacy_char_overlap(self) -> int:
        return max(0, int(self.settings.rag_chunk_overlap_tokens * 4))

    def _create_chunks(
        self,
        documents: list,
        max_chunk_size: int | None = None,
        overlap: int | None = None,
    ) -> list[str]:
        if max_chunk_size is None:
            max_chunk_size = self._legacy_char_chunk_size()
        if overlap is None:
            overlap = self._legacy_char_overlap()
        separators = ["\n\n", "\n", " ", ""]

        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=max_chunk_size,
            chunk_overlap=overlap,
            length_function=len,
            separators=separators,
        )
        split_docs = text_splitter.split_documents(documents)
        return [doc.page_content for doc in split_docs]

    def _process_excel_workbook(
        self,
        file_path: str,
        original_filename: str | None = None,
    ) -> list[dict[str, Any]]:
        """Convert an Excel workbook into text chunks without MinerU."""
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to parse .xlsx uploads") from exc

        formula_workbook = load_workbook(file_path, data_only=False, read_only=True)
        value_workbook = load_workbook(file_path, data_only=True, read_only=True)

        try:
            chunks_with_metadata: list[dict[str, Any]] = []
            value_sheets = {sheet.title: sheet for sheet in value_workbook.worksheets}

            for sheet_index, formula_sheet in enumerate(formula_workbook.worksheets):
                value_sheet = value_sheets.get(formula_sheet.title)
                rows = self._extract_excel_rows(formula_sheet, value_sheet)
                if not rows:
                    continue

                text = self._excel_rows_to_markdown(formula_sheet.title, rows)
                chunks_with_metadata.append(
                    {
                        "text": text,
                        "page_start": sheet_index,
                        "page_end": sheet_index,
                        "has_images": False,
                        "image_count": 0,
                        "has_tables": True,
                        "table_count": 1,
                        "sheet_name": formula_sheet.title,
                        "source": original_filename,
                    }
                )

            if chunks_with_metadata:
                return chunks_with_metadata

            return [
                {
                    "text": (
                        f"Workbook {original_filename or Path(file_path).name} "
                        "contains no non-empty sheets."
                    ),
                    "page_start": None,
                    "page_end": None,
                    "has_images": False,
                    "image_count": 0,
                    "has_tables": False,
                    "table_count": 0,
                    "source": original_filename,
                }
            ]
        finally:
            formula_workbook.close()
            value_workbook.close()

    def _extract_excel_rows(
        self, formula_sheet: Any, value_sheet: Any | None
    ) -> list[list[str]]:
        rows: list[list[str]] = []
        value_rows = value_sheet.iter_rows() if value_sheet is not None else None

        for formula_row in formula_sheet.iter_rows():
            value_row = next(value_rows, []) if value_rows is not None else []
            values: list[str] = []

            for index, formula_cell in enumerate(formula_row):
                value_cell = value_row[index] if index < len(value_row) else None
                display_value = (
                    value_cell.value
                    if value_cell is not None and value_cell.value is not None
                    else formula_cell.value
                )
                values.append(self._stringify_excel_cell(display_value))

            while values and not values[-1]:
                values.pop()

            if values and any(value.strip() for value in values):
                rows.append(values)

        return rows

    @classmethod
    def _excel_rows_to_markdown(cls, sheet_name: str, rows: list[list[str]]) -> str:
        column_count = max((len(row) for row in rows), default=0)
        if column_count == 0:
            return f"# Sheet: {sheet_name}"

        padded_rows = [row + [""] * (column_count - len(row)) for row in rows]
        header = [
            value if value else f"Column {index + 1}"
            for index, value in enumerate(padded_rows[0])
        ]
        data_rows = padded_rows[1:]

        lines = [
            f"# Sheet: {sheet_name}",
            "",
            "| " + " | ".join(cls._escape_markdown_table_cell(v) for v in header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        for row in data_rows:
            lines.append(
                "| " + " | ".join(cls._escape_markdown_table_cell(v) for v in row) + " |"
            )

        return "\n".join(lines).strip()

    @staticmethod
    def _stringify_excel_cell(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        return str(value)

    @staticmethod
    def _escape_markdown_table_cell(value: str) -> str:
        return (
            value.replace("\\", "\\\\")
            .replace("|", "\\|")
            .replace("\r", " ")
            .replace("\n", " ")
        )

    def _parse_content_list_json(self, content_list_path: Path) -> list[dict[str, Any]]:
        """Parse MinerU's content_list.json into structured content blocks."""
        with open(content_list_path, encoding="utf-8") as f:
            content_list = json.load(f)
        return content_list

    def _create_chunks_with_page_metadata(
        self,
        content_blocks: list[dict[str, Any]],
        page_to_images: dict[int, list[dict[str, Any]]],
        max_chunk_size: int | None = None,
    ) -> list[dict[str, Any]]:
        """Create chunks from content_list.json blocks while preserving page metadata."""
        if max_chunk_size is None:
            max_chunk_size = self._legacy_char_chunk_size()

        chunks_with_metadata: list[dict[str, Any]] = []
        current_chunk_text = ""
        current_page_start: int | None = None
        current_page_end: int | None = None
        current_images: list[dict[str, Any]] = []
        current_tables: list[dict[str, Any]] = []
        pages_in_current_chunk: set = set()

        def _finalize_chunk() -> None:
            nonlocal current_chunk_text, current_page_start, current_page_end
            nonlocal current_images, current_tables, pages_in_current_chunk

            if not current_chunk_text.strip():
                return

            chunk_images = current_images.copy()
            for page in pages_in_current_chunk:
                if page in page_to_images:
                    for img in page_to_images[page]:
                        if img not in chunk_images:
                            chunk_images.append(img)

            chunks_with_metadata.append(
                {
                    "text": current_chunk_text.strip(),
                    "page_start": current_page_start,
                    "page_end": current_page_end,
                    "has_images": bool(chunk_images),
                    "image_count": len(chunk_images),
                    "images": chunk_images,
                    "has_tables": bool(current_tables),
                    "table_count": len(current_tables),
                    "tables": current_tables,
                }
            )

            current_chunk_text = ""
            current_page_start = None
            current_page_end = None
            current_images = []
            current_tables = []
            pages_in_current_chunk = set()

        for block in content_blocks:
            page_idx = block.get("page_idx", 0)
            block_type = block.get("type", "text")
            bbox = block.get("bbox")
            text_level = block.get("text_level", 0)

            text = ""
            if block_type == "text":
                text = block.get("text", "")
                if text_level and text_level > 0:
                    heading_prefix = "#" * text_level + " "
                    text = heading_prefix + text
            elif block_type == "table":
                table_entry = {
                    "page": page_idx,
                    "bbox": bbox,
                    "caption": block.get("table_caption", []),
                    "footnote": block.get("table_footnote", []),
                    "body": block.get("table_body", ""),
                }
                current_tables.append(table_entry)
                captions = block.get("table_caption", [])
                footnotes = block.get("table_footnote", [])
                if captions:
                    text = "[Table: " + " ".join(captions) + "]"
                elif footnotes:
                    text = "[Table] " + " ".join(str(note) for note in footnotes if note)
                else:
                    text = "[Table]"
            elif block_type == "image":
                image_entry = {
                    "page": page_idx,
                    "bbox": bbox,
                    "path": block.get("img_path", ""),
                    "caption": block.get("image_caption", []),
                    "footnote": block.get("image_footnote", []),
                }
                current_images.append(image_entry)
                captions = block.get("image_caption", [])
                footnotes = block.get("image_footnote", [])
                if captions:
                    text = "[Image: " + " ".join(captions) + "]"
                elif footnotes:
                    text = "[Image] " + " ".join(str(note) for note in footnotes if note)
                else:
                    text = "[Image]"
            elif block_type == "equation":
                text = block.get("text", "")
                if not text:
                    text = "[Equation]"
            else:
                text = block.get("text", "")

            proposed_length = len(current_chunk_text) + len(text) + 1
            if current_chunk_text and proposed_length > max_chunk_size:
                _finalize_chunk()

            if current_page_start is None:
                current_page_start = page_idx
            current_page_end = page_idx
            pages_in_current_chunk.add(page_idx)

            if text:
                current_chunk_text += text + "\n"

        _finalize_chunk()

        logger.info(
            "Created %d page-aware chunks from %d content blocks",
            len(chunks_with_metadata),
            len(content_blocks),
        )

        return chunks_with_metadata

    # --- Filename normalization helpers ---

    @staticmethod
    def _normalize_filename_token(value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value or "")
        normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
        normalized = normalized.lower()
        normalized = re.sub(r"[^a-z0-9]+", "-", normalized)
        normalized = re.sub(r"-+", "-", normalized).strip("-")
        return normalized

    @staticmethod
    def _collapse_filename_token(value: str) -> str:
        normalized = DocumentParseService._normalize_filename_token(value)
        return normalized.replace("-", "")

    def _build_filename_aliases(
        self, sanitized_stem: str, original_filename: str | None = None
    ) -> list[str]:
        aliases: list[str] = []

        def _add_alias(value: str | None) -> None:
            if value is None:
                return
            candidate = value.strip()
            if candidate and candidate not in aliases:
                aliases.append(candidate)

        _add_alias(sanitized_stem)
        _add_alias(self._normalize_filename_token(sanitized_stem))

        if sanitized_stem and "_" in sanitized_stem:
            without_prefix = sanitized_stem.split("_", 1)[1]
            _add_alias(without_prefix)
            _add_alias(self._normalize_filename_token(without_prefix))

        if original_filename:
            original_stem = Path(original_filename).stem
            _add_alias(original_stem)
            _add_alias(self._normalize_filename_token(original_stem))
            original_suffix = Path(original_filename).suffix.lower()
        else:
            original_suffix = ""

        if original_suffix:
            existing_aliases = list(aliases)
            for alias in existing_aliases:
                alias_with_ext = f"{alias}{original_suffix}"
                if alias_with_ext not in aliases:
                    aliases.append(alias_with_ext)

        return aliases

    def _resolve_mineru_output_dir(
        self, output_dir: Path, filename_candidates: list[str]
    ) -> Path | None:
        if not output_dir.exists():
            return None

        ordered_candidates: list[str] = []
        for candidate in filename_candidates or []:
            if candidate and candidate not in ordered_candidates:
                ordered_candidates.append(candidate)

        if not ordered_candidates:
            return None

        for candidate in ordered_candidates:
            expected_dir = output_dir / candidate
            if expected_dir.exists():
                return expected_dir

        normalized_targets = list(
            filter(
                None,
                (
                    self._normalize_filename_token(candidate)
                    for candidate in ordered_candidates
                    if candidate
                ),
            )
        )
        normalized_target_set = set(normalized_targets)
        collapsed_target_set = set(
            filter(
                None,
                (
                    self._collapse_filename_token(candidate)
                    for candidate in ordered_candidates
                    if candidate
                ),
            )
        )

        normalized_matches: list[Path] = []
        child_dirs = [child for child in output_dir.iterdir() if child.is_dir()]

        for child in child_dirs:
            normalized_child = self._normalize_filename_token(child.name)
            collapsed_child = self._collapse_filename_token(child.name)

            def _matches_target(token: str | None, token_set: set[str]) -> bool:
                if not token or not token_set:
                    return False
                if token in token_set:
                    return True
                return any(
                    token.endswith(target)
                    or token.startswith(target)
                    or target.endswith(token)
                    or target.startswith(token)
                    for target in token_set
                )

            if _matches_target(normalized_child, normalized_target_set) or _matches_target(
                collapsed_child, collapsed_target_set
            ):
                normalized_matches.append(child)

        if normalized_matches:
            chosen = sorted(
                normalized_matches,
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )[0]
            return chosen

        if child_dirs:
            chosen = sorted(
                child_dirs,
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )[0]
            return chosen

        return None

    def _resolve_markdown_file(
        self, search_roots: list[Path], filename_candidates: list[str]
    ) -> Path:
        ordered_candidates: list[str] = []
        for candidate in filename_candidates or []:
            if candidate and candidate not in ordered_candidates:
                ordered_candidates.append(candidate)

        if not ordered_candidates:
            raise RuntimeError("MinerU output missing filename hints")

        def _is_markdown_file(path: Path) -> bool:
            suffixes = [suffix.lower() for suffix in path.suffixes]
            if not suffixes:
                lowered_name = path.name.lower()
                return (
                    lowered_name.endswith(".md")
                    or lowered_name.endswith(".markdown")
                    or lowered_name.endswith(".mdx")
                )
            return any(suffix in {".md", ".markdown", ".mdx"} for suffix in suffixes)

        markdown_suffixes = {".md", ".markdown", ".mdx"}

        def _search_in_root(root: Path) -> Path | None:
            if root is None or not root.exists():
                return None

            for candidate in ordered_candidates:
                exact_matches = sorted(
                    root.glob(f"**/{candidate}.md"),
                    key=lambda p: len(p.parts),
                )
                if exact_matches:
                    return exact_matches[0]

            all_markdown = sorted(
                (
                    path
                    for path in root.rglob("*")
                    if path.is_file()
                    and (
                        path.suffix.lower() in markdown_suffixes or _is_markdown_file(path)
                    )
                ),
                key=lambda p: (len(p.parts), p.stat().st_mtime),
            )

            normalized_map: dict[str, Path] = {}
            collapsed_map: dict[str, Path] = {}
            for path in all_markdown:
                normalized_name = self._normalize_filename_token(path.stem)
                if normalized_name and normalized_name not in normalized_map:
                    normalized_map[normalized_name] = path
                collapsed_name = self._collapse_filename_token(path.stem)
                if collapsed_name and collapsed_name not in collapsed_map:
                    collapsed_map[collapsed_name] = path

            for candidate in ordered_candidates:
                normalized_candidate = self._normalize_filename_token(candidate)
                if normalized_candidate and normalized_candidate in normalized_map:
                    return normalized_map[normalized_candidate]
                collapsed_candidate = self._collapse_filename_token(candidate)
                if collapsed_candidate and collapsed_candidate in collapsed_map:
                    return collapsed_map[collapsed_candidate]

            normalized_candidates = [
                self._normalize_filename_token(candidate)
                for candidate in ordered_candidates
                if candidate
            ]
            normalized_candidates = [t for t in normalized_candidates if t]
            collapsed_candidates = [
                self._collapse_filename_token(candidate)
                for candidate in ordered_candidates
                if candidate
            ]
            collapsed_candidates = [t for t in collapsed_candidates if t]

            for path in all_markdown:
                normalized_name = self._normalize_filename_token(path.stem) or ""
                collapsed_name = self._collapse_filename_token(path.stem)
                if not normalized_name and not collapsed_name:
                    continue
                for normalized_candidate in normalized_candidates:
                    if (
                        normalized_candidate in normalized_name
                        or normalized_name in normalized_candidate
                    ):
                        return path
                for collapsed_candidate in collapsed_candidates:
                    if (
                        collapsed_candidate in collapsed_name
                        or collapsed_name in collapsed_candidate
                    ):
                        return path

            if all_markdown:
                return all_markdown[0]

            return None

        unique_roots: list[Path] = []
        for root in search_roots or []:
            if root and root not in unique_roots and root.exists():
                unique_roots.append(root)

        default_output_root = Path("output")
        if default_output_root.exists() and default_output_root not in unique_roots:
            unique_roots.append(default_output_root)

        if not unique_roots:
            raise RuntimeError(
                f"MinerU output missing markdown file for {ordered_candidates[0]}"
            )

        searched_roots: list[Path] = []
        for root in unique_roots:
            searched_roots.append(root)
            result = _search_in_root(root)
            if result:
                if root != unique_roots[0]:
                    logger.warning(
                        "MinerU markdown resolved via fallback root %s for %s",
                        root,
                        ordered_candidates[0],
                    )
                return result

        raise RuntimeError(
            f"MinerU output missing markdown file for {ordered_candidates[0]}"
        )
